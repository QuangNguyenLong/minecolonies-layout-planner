#!/usr/bin/env python3
"""
blueprints_to_excel.py
Convert a blueprint-pack JSON (nested "categories" / "blueprints", each blueprint
with size {x, y, z}) into an Excel workbook for laying out a colony.

Output workbook tabs:
  1. Buildings        - flat list (group, name, width x, length z, height y, footprint, links)
  2. Building Palette - every building drawn at 1 cell = 1 block, ORIGINAL on the
                        left and ROTATED 90 deg right next to it (same row)
  3. Colony Map       - empty square-cell grid (ruler every 10, chunk line every 16)

Usage:
  pip install openpyxl
  python blueprints_to_excel.py pack.json
  python blueprints_to_excel.py pack.json -o my_layout.xlsx --group-depth 2 --map-size 200

Options:
  -o, --output       output .xlsx (default: <displayName>_layout.xlsx)
  --group-depth N    how many category levels form a color group (default 2)
                     e.g. depth 1 -> "agriculture", depth 2 -> "agriculture / husbandry"
  --max-width N      max palette width in columns before wrapping (default 150)
  --map-size N       Colony Map is N x N cells (default 160)
"""
import argparse, colorsys, json, re, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

FONT = "Arial"


# --------------------------------------------------------------------------- parse
def collect(node, path, out):
    """Recursively walk a node; append (category_path_tuple, name, x, y, z)."""
    for name, bp in (node.get("blueprints") or {}).items():
        size = bp.get("size")
        if not size or not all(k in size for k in "xyz"):
            print(f"  skipped (no size): {'/'.join(path + (name,))}", file=sys.stderr)
            continue
        out.append((path, name, int(size["x"]), int(size["y"]), int(size["z"])))
    for cname, child in (node.get("categories") or {}).items():
        collect(child, path + (cname,), out)


def load(json_path, depth):
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    raw = []
    collect(data, (), raw)
    rows = []  # (group, display_name, x, y, z)
    for path, name, x, y, z in raw:
        group = " / ".join(path[:depth]) if path else "root"
        rest = path[depth:]                      # deeper categories go into the name
        disp = ": ".join(list(rest) + [name]) if rest else name
        rows.append((group, disp, x, y, z))
    return data.get("displayName") or Path(json_path).stem, rows


# --------------------------------------------------------------------------- colors
def make_colors(groups):
    """Deterministic pastel color per group (hex, no #)."""
    cols, n = {}, max(len(groups), 1)
    for i, g in enumerate(groups):
        h = (i / n + (0.07 if i % 2 else 0)) % 1.0
        s = 0.45 if i % 2 else 0.60
        v = 0.97 if i % 2 else 0.90
        r, gg, b = colorsys.hsv_to_rgb(h, s, v)
        cols[g] = f"{int(r*255):02X}{int(gg*255):02X}{int(b*255):02X}"
    return cols


# --------------------------------------------------------------------------- build
def build(title, rows, out_path, max_w, map_n):
    groups = sorted({r[0] for r in rows}, key=lambda g: [r[0] for r in rows].index(g))
    COL = make_colors(groups)
    thick = Side(style="medium", color="000000")

    wb = Workbook()
    ws = wb.active
    ws.title = "Buildings"
    pal = wb.create_sheet("Building Palette")
    cm = wb.create_sheet("Colony Map")

    # ---- palette
    pal.sheet_format.defaultColWidth = 2.7
    pal.sheet_format.defaultRowHeight = 18
    for c in range(1, max_w + 4):
        pal.column_dimensions[L(c)].width = 2.7
    pal["B1"] = ("BUILDING PALETTE - 1 cell = 1 block. Each building shown twice: ORIGINAL (left, "
                 "width = x / length = z) and ROTATED 90 deg (right, marked 'rot').")
    pal["B1"].font = Font(name=FONT, bold=True, size=12)
    pal["B2"] = ("Copy a block (select its rectangle, Ctrl+C) and paste into 'Colony Map'. "
                 "Name sits in each block's top-left cell.")
    pal["B2"].font = Font(name=FONT, italic=True, size=9)

    def draw(g, name, r0, c0, w, h, rot):
        fill = PatternFill("solid", fgColor=COL[g])
        for r in range(r0, r0 + h):
            for c in range(c0, c0 + w):
                cell = pal.cell(row=r, column=c)
                cell.fill = fill
                cell.border = Border(left=thick if c == c0 else None,
                                     right=thick if c == c0 + w - 1 else None,
                                     top=thick if r == r0 else None,
                                     bottom=thick if r == r0 + h - 1 else None)
        t = pal.cell(row=r0, column=c0, value=name + (" (rot)" if rot else ""))
        t.font = Font(name=FONT, bold=True, size=8)
        t.alignment = Alignment(horizontal="left", vertical="center")
        if h >= 2:
            s = pal.cell(row=r0 + 1, column=c0, value=f"{w}x{h}")
            s.font = Font(name=FONT, size=8)
            s.alignment = Alignment(horizontal="left", vertical="center")

    addr, addr_rot = {}, {}
    row = 4
    for g in groups:
        items = [r for r in rows if r[0] == g]
        items.sort(key=lambda d: (-max(d[2], d[4]), -d[2], d[1]))
        pal.cell(row=row, column=2, value=g.upper()).font = Font(name=FONT, bold=True, size=11)
        for c in range(2, 2 + len(g) + 4):
            pal.cell(row=row, column=c).fill = PatternFill("solid", fgColor=COL[g])
        row += 2
        cx, shelf_h, shelf_top = 2, 0, row
        for (_, n, x, y, z) in items:
            unit_w = x + 1 + z
            if cx + unit_w > max_w and cx > 2:       # wrap to a new shelf
                shelf_top += shelf_h + 1
                cx, shelf_h = 2, 0
            draw(g, n, shelf_top, cx, x, z, False)
            addr[(g, n)] = f"{L(cx)}{shelf_top}"
            draw(g, n, shelf_top, cx + x + 1, z, x, True)
            addr_rot[(g, n)] = f"{L(cx + x + 1)}{shelf_top}"
            cx += unit_w + 3
            shelf_h = max(shelf_h, x, z)
        row = shelf_top + shelf_h + 3
    pal.freeze_panes = "A4"

    # ---- buildings list
    hdr = ["Group", "Building", "Width (x)", "Length (z)", "Height (y)",
           "Footprint (blocks)", "Palette cell (original)", "Palette cell (rotated)"]
    ws.append(hdr)
    for c in range(1, 9):
        h = ws.cell(row=1, column=c)
        h.font = Font(name=FONT, bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="305496")
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (g, n, x, y, z) in enumerate(rows, start=2):
        ws.append([g, n, x, z, y, f"=C{i}*D{i}", addr[(g, n)], addr_rot[(g, n)]])
        for c in range(1, 7):
            ws.cell(row=i, column=c).font = Font(name=FONT)
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=COL[g])
        for c in (3, 4, 5, 6):
            ws.cell(row=i, column=c).alignment = Alignment(horizontal="center")
        for c, ad in ((7, addr[(g, n)]), (8, addr_rot[(g, n)])):
            cell = ws.cell(row=i, column=c)
            # internal link: use `location`, NOT a "#Sheet!A1" string (that breaks in Excel)
            cell.hyperlink = Hyperlink(ref=cell.coordinate,
                                       location=f"'Building Palette'!{ad}", display=ad)
            cell.font = Font(name=FONT, color="0563C1", underline="single")
            cell.alignment = Alignment(horizontal="center")
    last = len(rows) + 1
    for col, w in zip("ABCDEFGH", [26, 38, 11, 11, 11, 18, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{last}"
    ws.cell(row=last + 2, column=1,
            value=f"Source: {title} blueprint pack JSON (size.x / size.y / size.z). "
                  "Width = x, Length = z, Height = y.").font = Font(name=FONT, italic=True, size=9)

    # ---- colony map
    cm.sheet_format.defaultColWidth = 2.7
    cm.sheet_format.defaultRowHeight = 18
    for c in range(1, map_n + 3):
        cm.column_dimensions[L(c)].width = 2.7
    cm.column_dimensions["A"].width = 4.5
    light = Side(style="thin", color="808080")
    for i in range(map_n):
        if i % 10 == 0:
            cm.cell(row=1, column=i + 2, value=i).font = Font(name=FONT, size=6, color="7F7F7F")
            cm.cell(row=i + 2, column=1, value=i).font = Font(name=FONT, size=6, color="7F7F7F")
        if i % 16 == 0:
            for r in range(2, map_n + 2):
                cm.cell(row=r, column=i + 2).border = Border(left=light)
    for i in range(0, map_n, 16):
        for c in range(2, map_n + 2):
            cell = cm.cell(row=i + 2, column=c)
            cell.border = Border(left=cell.border.left, top=light)
    cm.freeze_panes = "B2"

    pal.sheet_properties.tabColor = "ED7D31"
    cm.sheet_properties.tabColor = "70AD47"
    wb.save(out_path)


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Blueprint JSON -> colony layout Excel")
    ap.add_argument("json_file")
    ap.add_argument("-o", "--output")
    ap.add_argument("--group-depth", type=int, default=2)
    ap.add_argument("--max-width", type=int, default=150)
    ap.add_argument("--map-size", type=int, default=160)
    a = ap.parse_args()

    title, rows = load(a.json_file, a.group_depth)
    if not rows:
        sys.exit("No blueprints with a size found in that JSON.")
    widest = max(x + 1 + z for _, _, x, _, z in rows) + 3  # widest original+rotated pair
    max_w = max(a.max_width, widest)
    out = a.output or re.sub(r"[^\w]+", "_", title).strip("_").lower() + "_layout.xlsx"
    build(title, rows, out, max_w, a.map_size)
    print(f"{len(rows)} buildings -> {out}")


if __name__ == "__main__":
    main()